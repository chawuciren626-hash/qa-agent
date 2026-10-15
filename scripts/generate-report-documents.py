import json
import os
import re
import sys
import hashlib
from html import escape
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from report_delivery import markdown_body, markdown_blocks, export_browser_delivery

if len(sys.argv) not in (6, 7):
    raise SystemExit("用法: python generate-report-documents.py <report.md> <results.json> <bugs.json> <report.docx> <report.pdf> [quality-metrics.json]")

report_md, results_path, bugs_path, docx_path, pdf_path = map(os.path.abspath, sys.argv[1:6])
metrics_path = os.path.abspath(sys.argv[6]) if len(sys.argv) == 7 else None
with open(report_md, encoding="utf-8") as f:
    markdown = markdown_body(f.read())
with open(results_path, encoding="utf-8") as f:
    results = json.load(f)
with open(bugs_path, encoding="utf-8") as f:
    bugs = json.load(f)
def validate_report_results(metrics, results):
    execution = metrics.get("execution", {})
    if execution and (execution.get("status") != results.get("status") or any(
        execution.get(key, 0) != results.get("summary", {}).get(key, 0)
        for key in ("total", "passed", "failed", "skipped", "blocked")
    )):
        raise SystemExit("Results conflict with current metrics; use report/reportable-results.json.")
    state = metrics.get("trust", {}).get("execution_state")
    if state and state.get("invocation_id") != results.get("invocation_id"):
        raise SystemExit("Report results belong to another execution invocation.")

def validate_report_lifecycle(metrics, ledger_bytes):
    governance = metrics.get('governance')
    if not governance:
        raise SystemExit('Lifecycle metrics are required; rebuild quality metrics.')
    if governance.get('integrity_error'):
        raise SystemExit('Lifecycle evidence is invalid; formal report export blocked.')
    actual = hashlib.sha256(ledger_bytes).hexdigest() if ledger_bytes is not None else None
    if actual != governance.get('ledger_sha256'):
        raise SystemExit('Lifecycle changed; rebuild quality metrics before export.')


metrics = {}
if metrics_path:
    with open(metrics_path, encoding="utf-8") as f:
        metrics = json.load(f)
    current_state_path = os.path.join(os.path.dirname(metrics_path), '..', 'execution', 'execution-state.json')
    if os.path.isfile(current_state_path):
        with open(current_state_path, encoding='utf-8') as f:
            current_state = json.load(f)
        if current_state != metrics.get('trust', {}).get('execution_state'):
            raise SystemExit('Execution state changed; rebuild quality metrics before exporting.')
    validate_report_results(metrics, results)
    ledger_path = os.path.join(os.path.dirname(metrics_path), '..', 'defects', 'lifecycle.json')
    ledger_bytes = None
    if os.path.isfile(ledger_path):
        with open(ledger_path, 'rb') as ledger:
            ledger_bytes = ledger.read()
    validate_report_lifecycle(metrics, ledger_bytes)
else:
    current_state_path = os.path.join(os.path.dirname(results_path), '..', 'execution', 'execution-state.json')
    if os.path.isfile(current_state_path):
        raise SystemExit('Current execution state requires quality metrics for report export.')

def markdown_lines(text):
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("---") or line.startswith("| ") or line.startswith("|-"):
            continue
        if line.startswith("#"):
            yield "heading", re.sub(r"^#+\s*", "", line)
        elif re.match(r"^[-*]\s+", line):
            yield "bullet", re.sub(r"^[-*]\s+", "", line)
        elif re.match(r"^\d+\.\s+", line):
            yield "bullet", re.sub(r"^\d+\.\s+", "", line)
        elif not line.startswith(">"):
            yield "body", line

def clean_text(text):
    # Keep underscores in evidence paths; XML escaping belongs to the renderer.
    return re.sub(r"(`+|\*{1,3})", "", text)

summary = results.get("summary", {})
report_title = next((clean_text(line[2:].strip()) for line in markdown.splitlines() if line.startswith('# ')), '测试报告')
run_id = results.get("run_id", "未标识")
bug_items = bugs if isinstance(bugs, list) else bugs.get("bugs", [])
governance = metrics.get('governance', {})
if governance:
    bug_items = governance.get('current', [])
lifecycle_items = governance.get('current', []) + governance.get('history', [])
quality_gate = metrics.get("quality_gate", {})
trust = metrics.get("trust", {}).get("evidence", {})
gate_status = quality_gate.get("status")
historical_open = governance.get('open_historical')
history_text = (f"历史未关闭问题: {len(historical_open)}" if isinstance(historical_open, list)
                else "历史未关闭问题: 未评估") + "；fixed仅表示修复声明，不等于复测通过。"
gate_reasons = quality_gate.get("reasons", [])

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run(report_title)
run.font.color.rgb = RGBColor(0, 0, 0)
run.font.name = "Microsoft YaHei"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
doc.add_paragraph(f"运行编号: {run_id}    生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
if gate_status:
    doc.add_paragraph(f"质量门状态: {gate_status}")
if gate_reasons:
    doc.add_paragraph("质量门原因: " + ", ".join(gate_reasons))
if trust:
    doc.add_paragraph(f"失败证据: 可分析 {trust.get('triage_ready', 0)}，证据不足 {trust.get('insufficient', 0)}")
doc.add_heading("执行概览", level=1)
table = doc.add_table(rows=2, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
for cell, value in zip(table.rows[0].cells, ["总计", "通过", "失败", "跳过", "阻塞"]): cell.text = value
for cell, value in zip(table.rows[1].cells, [summary.get("total", 0), summary.get("passed", 0), summary.get("failed", 0), summary.get("skipped", 0), summary.get("blocked", 0)]): cell.text = str(value)
for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs: paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, text in markdown_blocks(markdown):
    if kind == 'table':
        width = max(len(row) for row in text)
        body_table = doc.add_table(rows=0, cols=width)
        body_table.style = 'Table Grid'
        for row_index, values in enumerate(text):
            cells = body_table.add_row().cells
            for i, value in enumerate(values): cells[i].text = clean_text(value)
            if row_index == 0:
                repeat = OxmlElement('w:tblHeader'); body_table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for cell in cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs: run.bold = True
        continue
    if text in (report_title, "测试报告", "执行概览"): continue
    if kind == "heading": doc.add_heading(clean_text(text), level=2)
    elif kind == "bullet": doc.add_paragraph(clean_text(text), style="List Bullet")
    else: doc.add_paragraph(clean_text(text))
doc.add_heading("缺陷汇总", level=1)
doc.add_paragraph(f"缺陷草稿条目: {len(bug_items)}（工程问题另见正文）")
if bug_items:
    bug_table = doc.add_table(rows=1, cols=4); bug_table.style = "Table Grid"
    for cell, value in zip(bug_table.rows[0].cells, ["编号", "分类", "关联用例", "结论"]): cell.text = value
    for index, item in enumerate(bug_items, 1):
        cells = bug_table.add_row().cells
        values = [item.get("defect_id", item.get("bug_id", f"BUG-{index:03d}")), item.get("category", item.get("classification", "needs_review")), ", ".join(item.get('related_cases', [])) or item.get("related_case", item.get("case_id", "")), item.get("summary", item.get("title", "待确认问题"))]
        for cell, value in zip(cells, values): cell.text = str(value)
doc.add_paragraph("报告依据本次运行中的测试结果和缺陷分析产物。")
if governance:
    doc.add_heading('问题生命周期', level=1)
    doc.add_paragraph(history_text)
for item in lifecycle_items:
    doc.add_paragraph(f"{item.get('defect_id', '')}: {item.get('lifecycle_status', '')}; origin={item.get('origin_invocation_id', results.get('invocation_id', ''))}")
os.makedirs(os.path.dirname(docx_path), exist_ok=True)
doc.save(docx_path)

font_path = os.environ.get('QA_REPORT_FONT')
if not font_path and os.name == 'nt':
    candidate = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts', 'simsun.ttc')
    if os.path.isfile(candidate):
        font_path = candidate
if font_path:
    pdf_font = 'ReportCJK'
    pdfmetrics.registerFont(TTFont(pdf_font, font_path, subfontIndex=0))
else:
    pdf_font = 'STSong-Light'
    pdfmetrics.registerFont(UnicodeCIDFont(pdf_font))
styles_pdf = getSampleStyleSheet()
title_style = ParagraphStyle("title", parent=styles_pdf["Title"], fontName=pdf_font, fontSize=18, leading=24, alignment=TA_CENTER)
heading_style = ParagraphStyle("heading", parent=styles_pdf["Heading2"], fontName=pdf_font, fontSize=13, leading=19, spaceBefore=10, spaceAfter=5, keepWithNext=True)
body_style = ParagraphStyle("body", parent=styles_pdf["BodyText"], fontName=pdf_font, fontSize=9.5, leading=15, wordWrap='CJK')
gate_text = f"质量门状态: {escape(gate_status)}" if gate_status else ''
if gate_reasons:
    gate_text += "<br/>质量门原因: " + ", ".join(gate_reasons)
if trust:
    gate_text += f"<br/>失败证据: 可分析 {trust.get('triage_ready', 0)}，证据不足 {trust.get('insufficient', 0)}"
story = [Paragraph(escape(report_title), title_style), Spacer(1, 5 * mm), Paragraph(f"运行编号: {run_id}<br/>生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}<br/>{gate_text}", body_style), Spacer(1, 4 * mm), Paragraph("执行概览", heading_style)]
overview = [["总计", "通过", "失败", "跳过", "阻塞"], [str(summary.get("total", 0)), str(summary.get("passed", 0)), str(summary.get("failed", 0)), str(summary.get("skipped", 0)), str(summary.get("blocked", 0))]]
table_pdf = Table(overview, colWidths=[34 * mm] * 5)
table_pdf.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), pdf_font), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C6D9")), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
story.append(table_pdf)
for kind, text in markdown_blocks(markdown):
    if kind == 'table':
        width = max(len(row) for row in text)
        table_style = ParagraphStyle('table-body', parent=body_style, fontSize=8, leading=12)
        head_style = ParagraphStyle('table-head', parent=table_style, textColor=colors.white)
        cells = [[Paragraph(escape(clean_text(v)), head_style if i == 0 else table_style) for v in row + [''] * (width - len(row))] for i, row in enumerate(text)]
        body_table = Table(cells, colWidths=[174 * mm / width] * width, repeatRows=1)
        body_table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#16324F')), ('VALIGN',(0,0),(-1,-1),'TOP'), ('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#DCE5EF')), ('TOPPADDING',(0,0),(-1,-1),5), ('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.extend([body_table, Spacer(1, 3 * mm)])
        continue
    if text in (report_title, "测试报告", "执行概览"): continue
    story.append(Paragraph(escape(clean_text(text)), heading_style if kind == "heading" else body_style))
story.extend([Paragraph("缺陷汇总", heading_style), Paragraph(f"缺陷草稿条目: {len(bug_items)}（工程问题另见正文）", body_style)])
if governance:
    story.extend([Paragraph('问题生命周期', heading_style), Paragraph(history_text, body_style)])
for item in lifecycle_items:
    story.append(Paragraph(escape(f"{item.get('defect_id', '')}: {item.get('lifecycle_status', '')}; origin={item.get('origin_invocation_id', results.get('invocation_id', ''))}"), body_style))
SimpleDocTemplate(pdf_path, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm).build(story)
export_browser_delivery(os.path.dirname(docx_path), report_title, markdown, results, bug_items, metrics)
print(f"已生成: {docx_path}\n已生成: {pdf_path}")
